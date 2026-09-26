from datetime import date
from io import BytesIO
import re

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .directory_services import filtered_student_rows
from .models import ActivityLog, Course, CourseInstructor, CourseSession, Instructor, Registration, Student, TrainingRecord
from .report_export import workbook_bytes


class ReviewCorrectionsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.admin = User.objects.create_user('correction-admin', password='Review-Only-Secure-2026!', is_superuser=True, is_staff=True)
        cls.user = User.objects.create_user('correction-instructor')
        cls.instructor = Instructor.objects.create(user=cls.user, name_english='Instructor With A Long Full Name', phone='+971 50 123 4567', email='instructor@example.com')
        cls.course = Course.objects.create(title_english='Correction course')
        cls.session = CourseSession.objects.create(course=cls.course, start_date=date(2026,9,3), end_date=date(2026,9,4), status='registration_closed')
        CourseInstructor.objects.create(session=cls.session, instructor=cls.instructor)
        cls.student = Student.objects.create(name_english='Student Full Name', name_arabic='اسم الطالب الكامل', eid='784199012345671', identity_status='verified')
        cls.record = TrainingRecord.objects.create(student=cls.student, session=cls.session, result='pass')

    def setUp(self):
        self.client.force_login(self.admin)

    def test_all_changed_pages_render_with_real_context(self):
        for name in ('dashboard', 'my_courses', 'students', 'courses_sessions', 'staff', 'inventory', 'reports'):
            with self.subTest(page=name):
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, 'corrections-20260923')

    def test_job_card_opens_full_workspace_and_preserves_course_controls(self):
        url = reverse('instructor_course_workspace', args=[self.session.public_id])
        self.assertContains(self.client.get(reverse('my_courses')), url)
        self.client.force_login(self.user)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        for name in ('course_roster_save', 'course_roster_save_download', 'stamped_lists', 'course_inventory_save', 'registration_qr'):
            self.assertContains(response, reverse(name, args=[self.session.public_id]))
        self.assertContains(response, 'Material Used')
        self.assertFalse(response.context['course_read_only'])

    def test_data_quality_navigation_removed_and_review_report_remains(self):
        response = self.client.get(reverse('reports'), {'report':'duplicates'})
        nav = response.content.decode().split('aria-label="Main navigation"',1)[1].split('</nav>',1)[0]
        self.assertNotIn('Data Quality', nav)
        self.assertContains(response, 'Duplicates and review')
        self.assertEqual(response.context['report_view'], 'duplicates')

    def test_export_obeys_course_and_calendar_filters(self):
        other = Course.objects.create(title_english='Outside filter')
        other_session = CourseSession.objects.create(course=other, start_date=date(2025,1,1), status='completed')
        TrainingRecord.objects.create(student=self.student, session=other_session, result='fail')
        params={'report':'student_pass', 'course':self.course.pk, 'date_from':'2026-09-01', 'date_to':'2026-09-30', 'export':'xlsx'}
        response=self.client.get(reverse('reports'),params)
        self.assertEqual(response.status_code,200)
        self.assertIn('.xlsx',response['Content-Disposition'])
        sheet=load_workbook(BytesIO(response.content))['Report']
        self.assertEqual(sheet.max_row,3)
        self.assertEqual(sheet['A3'].value,self.student.name_english)
        self.assertEqual(sheet['C3'].data_type,'s')
        self.assertEqual(sheet['F3'].value.date(),self.session.start_date)
        self.assertEqual(sheet['H3'].value,'Pass')
        params['date_from']='2027-01-01';params['date_to']='2027-12-31'
        empty=load_workbook(BytesIO(self.client.get(reverse('reports'),params).content))['Report']
        self.assertEqual(empty.max_row,2)

    def test_each_report_downloads_a_valid_workbook(self):
        from .report_views import REPORT_LABELS
        for key in REPORT_LABELS:
            with self.subTest(report=key):
                response=self.client.get(reverse('reports'),{'report':key,'export':'xlsx'})
                self.assertEqual(response.status_code,200)
                self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames,['Report','Filters'])

    def test_activity_export_is_not_truncated_to_preview_limit(self):
        ActivityLog.objects.bulk_create([ActivityLog(actor=self.admin,action='update',object_type='Example',object_id=str(i),description='Export row') for i in range(260)])
        response=self.client.get(reverse('reports'),{'report':'activity','export':'xlsx'})
        self.assertEqual(load_workbook(BytesIO(response.content))['Report'].max_row,262)

    def test_excel_does_not_execute_names_as_formulas(self):
        value=workbook_bytes('Test',['Name','Emirates ID','Date','Number'],[['=1+1','784199012345671',date(2026,9,23),4]])
        sheet=load_workbook(BytesIO(value))['Report']
        self.assertEqual(sheet['A3'].data_type,'s')
        self.assertEqual(sheet['A3'].value,'=1+1')
        self.assertEqual(sheet['D3'].value,4)

    def test_export_requires_report_permission(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('reports'),{'export':'xlsx'}).status_code,403)

    def test_staff_filters_work_in_the_separate_contact_columns(self):
        response=self.client.get(reverse('staff'),{'phone':'123 4567','email':'instructor@'})
        self.assertEqual([r['staff'].pk for r in response.context['rows']],[self.instructor.pk])
        self.assertContains(response,'Mobile number')
        self.assertContains(response,'data-edit-staff')

    def test_student_duplicate_review_is_visible_without_changing_identity(self):
        args=dict(student=self.student,requested_session=self.session,submitted_name_english='Student Full Name',submitted_name_arabic='اسم الطالب الكامل',eid_raw=self.student.eid,eid_normalized=self.student.eid,email_raw='student@example.com')
        Registration.objects.create(**args)
        Registration.objects.create(**args)
        rows=filtered_student_rows({'duplicates':'yes','identity_status':'review'})
        self.assertTrue(any(r['student'].pk==self.student.pk and r['duplicate_review'] for r in rows))
        self.student.refresh_from_db();self.record.refresh_from_db()
        self.assertEqual(self.student.identity_status,'verified')
        self.assertEqual(self.record.result,'pass')

    def test_login_works_with_csrf_and_rejects_missing_token(self):
        client=Client(enforce_csrf_checks=True)
        response=client.get(reverse('login'))
        token=re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"',response.content.decode()).group(1)
        credentials={'username':'correction-admin','password':'Review-Only-Secure-2026!'}
        self.assertEqual(client.post(reverse('login'),credentials).status_code,403)
        response=client.post(reverse('login'),{**credentials,'csrfmiddlewaretoken':token})
        self.assertEqual(response.status_code,302)
        self.assertEqual(int(client.session['_auth_user_id']),self.admin.pk)
