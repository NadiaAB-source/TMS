from io import BytesIO, StringIO
from pathlib import Path
import tempfile
import uuid
from datetime import date

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import course_workflow_views, stamped_list_views
from .course_services import (
    allocate_course_sequence_number,
    day_one_list_filename,
    day_one_students_upload_filename,
    enable_course_sequence_numbering,
)
from .identity import format_emirates_id, masked_emirates_id
from .j35_services import recommendation_cards
from .models import (
    ActivityLog,
    Camp,
    Course,
    CourseSequenceCounter,
    CourseInstructor,
    CourseInstructorInventoryUsage,
    CourseSession,
    CourseSessionProposal,
    CampContact,
    Instructor,
    InstructorAllocation,
    InstructorInventoryBalance,
    InstructorInventoryMovement,
    InstructorRole,
    J35TrainingNeed,
    InventoryItem,
    Notification,
    Registration,
    SourceFile,
    SourceRecord,
    StampedListArchive,
    Student,
    Team,
    TrainingRecord,
    UserSecurityProfile,
)
from .access import can_manage_j35_planner
from .staff_access_models import StaffPermissionOverride
from .staff_groups import (
    J35_PLANNER_GROUP,
    OPERATIONS_ADMIN_GROUP,
    ensure_predefined_groups,
    synchronise_staff_groups,
)


@override_settings(
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    DEFAULT_FROM_EMAIL="test@iqarus.example",
)
class SimpleCourseWorkflowTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="lead-admin",
            password="test-password",
            is_staff=True,
            is_superuser=True,
        )
        self.instructor_user = User.objects.create_user(
            username="course-instructor",
            email="instructor@example.com",
            password="test-password",
        )
        lead_role = InstructorRole.objects.create(name="LI")
        self.instructor = Instructor.objects.create(
            user=self.instructor_user,
            name_english="COURSE INSTRUCTOR",
            email="instructor@example.com",
        )
        self.instructor.roles.add(lead_role)
        self.unassigned_user = User.objects.create_user(
            username="unassigned-instructor",
            email="unassigned@example.com",
            password="test-password",
        )
        instructor_role = InstructorRole.objects.create(name="INS")
        self.unassigned_instructor = Instructor.objects.create(
            user=self.unassigned_user,
            name_english="UNASSIGNED INSTRUCTOR",
            email="unassigned@example.com",
        )
        self.unassigned_instructor.roles.add(instructor_role)
        self.course = Course.objects.create(
            code="ASM",
            title_english="ASM - TCCC",
        )
        self.camp = Camp.objects.create(name="Test Camp")
        self.tempdir = tempfile.TemporaryDirectory()
        self.original_roster_root = course_workflow_views.ROSTER_ROOT
        course_workflow_views.ROSTER_ROOT = Path(self.tempdir.name)
        self.original_stamped_root = stamped_list_views.STAMPED_LISTS_ROOT
        stamped_list_views.STAMPED_LISTS_ROOT = Path(self.tempdir.name) / "stamped"

    def tearDown(self):
        course_workflow_views.ROSTER_ROOT = self.original_roster_root
        stamped_list_views.STAMPED_LISTS_ROOT = self.original_stamped_root
        self.tempdir.cleanup()

    def _create_course(self, instructor_ids=None):
        self.client.force_login(self.admin)
        instructor_ids = instructor_ids or [self.instructor.id]
        response = self.client.post(
            reverse("course_create"),
            {
                "course": self.course.id,
                "camp": self.camp.id,
                "start_date": "2026-09-01",
                "end_date": "2026-09-02",
                "capacity": "20",
                "instructor_student_ratio": "10",
                "service_branch": "PG",
                "instructors": instructor_ids,
                "poc_name": "Camp Contact",
                "poc_contact_number": "+971500000000",
                "poc_location_url": "https://example.com/location",
            },
        )
        self.assertEqual(response.status_code, 302)
        return CourseSession.objects.get(course=self.course)

    def _registration(
        self,
        session,
        *,
        eid="784198012345678",
        name_english="TEST STUDENT",
        name_arabic="طالب تجريبي",
        email="student@example.com",
    ):
        return Registration.objects.create(
            requested_session=session,
            submitted_name_english=name_english,
            submitted_name_arabic=name_arabic,
            eid_raw=eid,
            eid_normalized=eid,
            email_raw=email,
            submitted_unit="Unit",
            status=Registration.Status.PENDING,
        )

    def _history_source_file(self):
        return SourceFile.objects.create(
            filename="historical-evidence.xlsx",
            sha256=uuid.uuid4().hex.ljust(64, "0"),
            records_found=1,
            imported=True,
        )

    def _history_record(
        self,
        source_file,
        *,
        record_type,
        eid,
        raw_payload,
        import_status=SourceRecord.ImportStatus.STAGED,
        linked_instructor=None,
        linked_registration=None,
        linked_training_record=None,
    ):
        return SourceRecord.objects.create(
            source_record_id=uuid.uuid4(),
            source_file=source_file,
            worksheet_name="History",
            source_row_number=SourceRecord.objects.count() + 2,
            record_type=record_type,
            raw_payload=raw_payload,
            derived_search_values={
                "normalized_emirates_id": eid,
            },
            import_status=import_status,
            linked_instructor=linked_instructor,
            linked_registration=linked_registration,
            linked_training_record=linked_training_record,
        )

    def _history_proposal(
        self,
        *records,
        start_date,
        course=None,
        instructor_values=None,
    ):
        proposal = CourseSessionProposal.objects.create(
            candidate_id=uuid.uuid4(),
            proposal_source=(
                CourseSessionProposal.ProposalSource.REGISTRATION
            ),
            proposed_course=course,
            proposed_camp=self.camp,
            start_date=start_date,
            end_date=start_date,
            camp_values=[self.camp.name],
            unit_values=["PG"],
            instructor_values=instructor_values or [],
            evidence_records=len(records),
            student_id_values=len(records),
        )
        proposal.source_records.add(*records)
        return proposal

    def test_emirates_id_display_always_uses_dashes(self):
        raw = "784198012345678"
        expected = "784-1980-1234567-8"
        self.assertEqual(format_emirates_id(raw), expected)
        self.assertEqual(format_emirates_id(expected), expected)
        self.assertEqual(
            masked_emirates_id(raw),
            "•••-••••-••••567-8",
        )

    def test_create_course_uses_assignment_focused_layout(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("course_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Course details")
        self.assertContains(response, "Service Branch")
        self.assertContains(response, "Point of Contact")
        self.assertContains(response, "Assigned instructors")
        self.assertContains(response, 'id="instructor-search"')
        self.assertContains(response, 'id="selected-instructor-count"')
        self.assertNotContains(response, "Lead Instructor")

    def test_course_creation_emails_assignment_and_creates_qr_token(self):
        session = self._create_course()
        assignment = CourseInstructor.objects.get(session=session)
        self.assertEqual(
            assignment.assignment_role,
            CourseInstructor.AssignmentRole.INSTRUCTOR,
        )
        self.assertIsNotNone(assignment.notified_at)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Camp Contact", mail.outbox[0].body)
        self.assertIn("Service Branch: PG", mail.outbox[0].body)
        self.assertIn("https://example.com/location", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["instructor@example.com"])
        self.assertFalse(
            CourseInstructor.objects.filter(
                session=session,
                instructor=self.unassigned_instructor,
            ).exists()
        )
        self.assertTrue(session.public_registration_token)
        self.assertFalse(session.registration_published)

        self.client.force_login(self.unassigned_user)
        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 403)

    def test_course_can_have_multiple_selected_instructors(self):
        session = self._create_course(
            [self.instructor.id, self.unassigned_instructor.id]
        )
        self.assertEqual(
            CourseInstructor.objects.filter(session=session).count(),
            2,
        )
        self.assertEqual(len(mail.outbox), 2)
        self.assertEqual(
            {message.to[0] for message in mail.outbox},
            {"instructor@example.com", "unassigned@example.com"},
        )

    def test_other_service_branch_requires_and_saves_custom_name(self):
        self.client.force_login(self.admin)
        course_data = {
            "course": self.course.id,
            "camp": self.camp.id,
            "start_date": "2026-09-01",
            "end_date": "2026-09-02",
            "capacity": "20",
            "instructor_student_ratio": "10",
            "service_branch": "OTHER",
            "instructors": [self.instructor.id],
            "poc_name": "Camp Contact",
            "poc_contact_number": "+971500000000",
            "poc_location_url": "https://example.com/location",
        }
        response = self.client.post(reverse("course_create"), course_data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter the service branch.")

        course_data["service_branch_other"] = "Coast Guard"
        response = self.client.post(reverse("course_create"), course_data)
        self.assertEqual(response.status_code, 302)
        session = CourseSession.objects.get(course=self.course)
        self.assertEqual(session.service_branch, "OTHER")
        self.assertEqual(session.service_branch_other, "Coast Guard")
        self.assertEqual(session.service_branch_name, "Coast Guard")
        self.assertIn("Service Branch: Coast Guard", mail.outbox[0].body)

    def test_authorised_user_edits_course_and_only_new_assignment_is_emailed(self):
        session = self._create_course()
        second_camp = Camp.objects.create(name="Second Camp")
        mail.outbox.clear()

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("course_edit", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Save Course Changes")

        response = self.client.post(
            reverse("course_edit", args=[session.public_id]),
            {
                "course": self.course.id,
                "camp": second_camp.id,
                "start_date": "2026-09-03",
                "end_date": "2026-09-04",
                "capacity": "30",
                "instructor_student_ratio": "15",
                "service_branch": "NAVY",
                "instructors": [self.unassigned_instructor.id],
                "poc_name": "Updated Contact",
                "poc_contact_number": "+971511111111",
                "poc_location_url": "https://example.com/updated-location",
            },
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.camp, second_camp)
        self.assertEqual(session.capacity, 30)
        self.assertEqual(session.service_branch, "NAVY")
        self.assertEqual(session.poc_name, "Updated Contact")
        self.assertTrue(session.reference_code.endswith("SECOND-CAMP"))
        self.assertFalse(
            CourseInstructor.objects.filter(
                session=session,
                instructor=self.instructor,
            ).exists()
        )
        new_assignment = CourseInstructor.objects.get(
            session=session,
            instructor=self.unassigned_instructor,
        )
        self.assertIsNotNone(new_assignment.notified_at)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["unassigned@example.com"])
        self.assertIn("Updated Contact", mail.outbox[0].body)
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.UPDATE,
            object_id=str(session.public_id),
            description="Course details or instructor assignments updated.",
        )
        self.assertEqual(
            activity.details["added_instructor_ids"],
            [self.unassigned_instructor.id],
        )
        self.assertEqual(
            activity.details["removed_instructor_ids"],
            [self.instructor.id],
        )

    def test_course_cancel_preserves_course_and_writes_activity_log(self):
        session = self._create_course()
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_cancel", args=[session.public_id]),
            {"reason": "Camp requested cancellation"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.status, CourseSession.Status.CANCELLED)
        self.assertFalse(session.registration_published)
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.CANCEL,
            object_id=str(session.public_id),
        )
        self.assertEqual(
            activity.details["reason"],
            "Camp requested cancellation",
        )

    def test_empty_course_can_be_deleted_but_activity_remains(self):
        session = self._create_course()
        public_id = str(session.public_id)
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_delete", args=[session.public_id]),
            {"confirm": "delete", "reason": "Created by mistake"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(CourseSession.objects.filter(public_id=public_id).exists())
        activity = ActivityLog.objects.get(
            action=ActivityLog.Action.DELETE,
            object_id=public_id,
        )
        self.assertEqual(activity.details["reason"], "Created by mistake")

    def test_course_with_registration_cannot_be_deleted(self):
        session = self._create_course()
        self._registration(session)
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("course_delete", args=[session.public_id]),
            {"confirm": "delete"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(CourseSession.objects.filter(pk=session.pk).exists())
        self.assertFalse(
            ActivityLog.objects.filter(
                action=ActivityLog.Action.DELETE,
                object_id=str(session.public_id),
            ).exists()
        )

    def test_inventory_is_held_and_reported_per_instructor(self):
        item = InventoryItem.objects.create(
            name="Training Kit",
            category=InventoryItem.Category.CONSUMABLE,
            unit="kit",
            quantity_on_hand=20,
        )
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("inventory"),
            {
                "action": "allocate",
                "instructor_id": self.unassigned_instructor.id,
                "item_id": item.id,
                "movement_type": "issue",
                "quantity": "8",
            },
        )
        self.assertEqual(response.status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.quantity_on_hand, 12)
        balance = InstructorInventoryBalance.objects.get(
            instructor=self.unassigned_instructor,
            item=item,
        )
        self.assertEqual(balance.quantity_on_hand, 8)

        session = self._create_course(
            [self.instructor.id, self.unassigned_instructor.id]
        )
        self.client.force_login(self.unassigned_user)
        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.unassigned_instructor.id,
                f"used_{item.id}": "5",
                f"consumed_{item.id}": "2",
                f"deteriorated_{item.id}": "1",
                f"inventory_notes_{item.id}": "Course use",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 5)
        usage = CourseInstructorInventoryUsage.objects.get(
            session=session,
            instructor=self.unassigned_instructor,
            item=item,
        )
        self.assertEqual(usage.quantity_used, 5)
        self.assertEqual(usage.quantity_consumed, 2)
        self.assertEqual(usage.quantity_deteriorated, 1)
        movement = InstructorInventoryMovement.objects.filter(
            session=session,
            instructor=self.unassigned_instructor,
            item=item,
        ).latest("created_at")
        self.assertEqual(movement.quantity_change, -3)
        self.assertEqual(movement.balance_after, 5)

        # Saving the same report again must not deduct the balance twice.
        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.unassigned_instructor.id,
                f"used_{item.id}": "5",
                f"consumed_{item.id}": "2",
                f"deteriorated_{item.id}": "1",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 5)

        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {"instructor_id": self.instructor.id},
        )
        self.assertEqual(response.status_code, 403)

    def test_material_used_layout_saves_one_used_up_value_per_instructor(self):
        item = InventoryItem.objects.create(
            name="Pressure Bandage",
            category=InventoryItem.Category.CONSUMABLE,
            unit="each",
            quantity_on_hand=0,
        )
        balance = InstructorInventoryBalance.objects.create(
            instructor=self.instructor,
            item=item,
            quantity_on_hand=20,
            updated_by=self.admin,
        )
        session = self._create_course()
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Material Used")
        self.assertContains(response, "Balance Before Course")
        self.assertContains(response, "Used Up in Course")
        self.assertContains(response, "Balance After Course")
        self.assertContains(response, f'name="used_up_{item.id}"')
        self.assertNotContains(response, f'name="consumed_{item.id}"')
        self.assertNotContains(response, f'name="deteriorated_{item.id}"')

        response = self.client.post(
            reverse("course_inventory_save", args=[session.public_id]),
            {
                "instructor_id": self.instructor.id,
                f"used_up_{item.id}": "5",
                f"inventory_notes_{item.id}": "5 damaged",
            },
        )
        self.assertEqual(response.status_code, 302)
        balance.refresh_from_db()
        self.assertEqual(balance.quantity_on_hand, 15)
        usage = CourseInstructorInventoryUsage.objects.get(
            session=session,
            instructor=self.instructor,
            item=item,
        )
        self.assertEqual(usage.quantity_used, 5)
        self.assertEqual(usage.quantity_consumed, 5)
        self.assertEqual(usage.quantity_deteriorated, 0)

    def test_inventory_records_use_tabs_filters_and_25_row_pages(self):
        for number in range(26):
            item = InventoryItem.objects.create(
                name=f"Filter Item {number:02d}",
                category=InventoryItem.Category.EQUIPMENT,
                unit="item",
                quantity_on_hand=0,
            )
            InstructorInventoryBalance.objects.create(
                instructor=self.instructor,
                item=item,
                quantity_on_hand=number + 1,
                updated_by=self.admin,
            )
            InstructorInventoryMovement.objects.create(
                instructor=self.instructor,
                item=item,
                movement_type=InstructorInventoryMovement.MovementType.ISSUE,
                quantity_change=number + 1,
                balance_after=number + 1,
                recorded_by=self.admin,
            )

        self.client.force_login(self.admin)
        response = self.client.get(reverse("inventory"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "portal/js/scroll_position.")
        self.assertContains(response, 'id="inventory-balances-panel"')
        self.assertNotContains(response, 'id="inventory-history-panel"')
        self.assertEqual(len(response.context["balance_page"]), 25)
        self.assertTrue(response.context["balance_page"].has_next())

        response = self.client.get(
            reverse("inventory"),
            {"view": "balances", "q": "Filter Item 25"},
        )
        self.assertEqual(response.context["balance_page"].paginator.count, 1)

        response = self.client.get(
            reverse("inventory"),
            {"view": "history"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="inventory-history-panel"')
        self.assertNotContains(response, 'id="inventory-balances-panel"')
        self.assertEqual(len(response.context["movement_page"]), 25)
        self.assertTrue(response.context["movement_page"].has_next())

        response = self.client.get(
            reverse("inventory"),
            {"view": "history", "page": "2"},
        )
        self.assertEqual(len(response.context["movement_page"]), 1)

    def test_assigned_instructor_controls_shared_course_registration(self):
        session = self._create_course()
        self.client.force_login(self.instructor_user)
        response = self.client.post(
            reverse("registration_qr", args=[session.public_id]),
            {"action": "open"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertTrue(session.registration_published)
        response = self.client.post(
            reverse("registration_qr", args=[session.public_id]),
            {"action": "close"},
        )
        self.assertEqual(response.status_code, 302)
        session.refresh_from_db()
        self.assertFalse(session.registration_published)

    def test_single_roster_saves_independent_tags_and_exact_template(self):
        session = self._create_course()
        registration = self._registration(session)
        self._registration(
            session,
            eid="784198112345679",
            name_english="NOT SELECTED STUDENT",
            name_arabic="طالب غير مختار",
            email="not-selected@example.com",
        )
        self.client.force_login(self.instructor_user)
        response = self.client.post(
            reverse("course_roster_save_download", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"hp_{registration.public_id}": "yes",
                f"ttt_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "passed",
                f"remarks_{registration.public_id}": "Clear practical result",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            "attachment; filename*=UTF-8''01-02_September_COURSE_Test_Camp.xlsx",
        )
        registration.refresh_from_db()
        self.assertTrue(registration.selected_for_roster)
        self.assertTrue(registration.is_hp)
        self.assertTrue(registration.is_ttt)
        self.assertEqual(registration.assessment_status, "passed")
        self.assertEqual(TrainingRecord.objects.count(), 1)
        record = TrainingRecord.objects.get()
        self.assertTrue(record.is_hp)
        self.assertTrue(record.is_ttt)

        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["List 1-1"])
        sheet = workbook.active
        self.assertEqual(
            [sheet.cell(3, column).value for column in range(1, 9)],
            [
                "#",
                "Email",
                "EID / رقم الهوية",
                "اسم الطالب",
                "Student Name",
                "Day 1",
                "Day 2",
                "Status",
            ],
        )
        self.assertEqual(sheet["B4"].value, "student@example.com")
        self.assertEqual(sheet["C4"].value, "784-1980-1234567-8")
        self.assertEqual(sheet["D4"].value, "طالب تجريبي")
        self.assertEqual(sheet["E4"].value, "TEST STUDENT")
        self.assertIsNone(sheet["H4"].value)
        self.assertEqual(sheet["E5"].value, "Status: Pass, Fail, D (Duplicate), TTT (Train the Trainer)")
        self.assertIsNone(sheet["A5"].value)
        self.assertTrue(str(sheet.print_area).endswith("$H$7"))
        self.assertNotIn(
            0,
            [
                cell.value
                for row in sheet.iter_rows(min_row=1, max_row=7)
                for cell in row
            ],
        )
        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertContains(response, "784-1980-1234567-8")
        self.assertNotContains(response, "Review required")

    def test_assigned_instructor_can_correct_roster_identity_details(self):
        session = self._create_course()
        registration = self._registration(session)
        self.client.force_login(self.instructor_user)

        response = self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "passed",
                f"email_{registration.public_id}": "Corrected.Email@EXAMPLE.COM",
                f"eid_{registration.public_id}": "784-1980-1234567-8",
                f"name_english_{registration.public_id}": "Corrected Student",
                f"name_arabic_{registration.public_id}": "طالب مصحح",
            },
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertEqual(registration.email_raw, "corrected.email@example.com")
        self.assertEqual(registration.eid_normalized, "784198012345678")
        self.assertEqual(registration.submitted_name_english, "CORRECTED STUDENT")
        self.assertEqual(registration.submitted_name_arabic, "طالب مصحح")
        self.assertIsNotNone(registration.student_id)
        self.assertEqual(registration.student.email, "corrected.email@example.com")
        self.assertEqual(registration.student.name_english, "CORRECTED STUDENT")
        activity = ActivityLog.objects.filter(
            object_type="CourseSession",
            object_id=str(session.public_id),
        ).latest("id")
        self.assertEqual(activity.details["edited_student_details"], 1)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertContains(response, f'name="email_{registration.public_id}"')
        self.assertContains(response, "Correct student details, record attendance")

    def test_course_sequence_is_dormant_until_the_post_reset_import(self):
        counter, _ = CourseSequenceCounter.objects.get_or_create(
            key="course_session",
            defaults={"next_number": 1},
        )
        self.assertFalse(counter.enabled)
        self.assertIsNone(allocate_course_sequence_number())

        enable_course_sequence_numbering()
        self.assertEqual(allocate_course_sequence_number(), 1)
        self.assertEqual(allocate_course_sequence_number(), 2)

    def test_stamped_list_is_the_only_document_workflow(self):
        session = self._create_course()
        self.client.force_login(self.instructor_user)
        url = reverse("stamped_lists", args=[session.public_id])

        response = self.client.post(
            url,
            {
                "document_type": "day_two_list",
                "document": SimpleUploadedFile(
                    "legacy.pdf",
                    b"%PDF-1.7\nlegacy",
                    content_type="application/pdf",
                ),
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(StampedListArchive.objects.count(), 0)

        response = self.client.post(
            url,
            {
                "document": SimpleUploadedFile(
                    "stamped.pdf",
                    b"%PDF-1.7\nstamped",
                    content_type="application/pdf",
                ),
            },
        )
        self.assertEqual(response.status_code, 302)
        document = StampedListArchive.objects.get()
        self.assertEqual(
            document.document_type,
            StampedListArchive.DocumentType.STAMPED_LIST,
        )

        response = self.client.get(url)
        self.assertContains(response, "Upload Stamped List")
        self.assertNotContains(response, "Day 2 List")

    def test_workspace_bulk_controls_are_in_their_table_columns(self):
        session = self._create_course()
        self._registration(session)
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="column-control">Select')
        self.assertContains(response, 'id="hp-all"')
        self.assertContains(response, 'id="ttt-all"')
        self.assertContains(response, 'id="pass-all"')
        self.assertContains(response, 'id="fail-all"')
        self.assertContains(response, 'class="student-row')
        self.assertContains(response, 'title="Session 1 attendance">S1')
        self.assertContains(response, 'title="Session 2 attendance">S2')
        self.assertNotContains(response, 'class="toolbar"')

    def test_dashboard_and_course_pages_use_reference_layout(self):
        session = self._create_course()
        registration = self._registration(session)
        registration.selected_for_roster = True
        registration.assessment_status = "passed"
        registration.is_hp = True
        registration.save(
            update_fields=[
                "selected_for_roster",
                "assessment_status",
                "is_hp",
                "updated_at",
            ]
        )
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="dashboard-filters"')
        self.assertContains(response, 'class="summary-counter-card"', count=3)
        self.assertContains(response, 'class="dashboard-chart"')
        self.assertContains(response, 'class="dashboard-overview"')
        self.assertContains(response, 'name="all_history"')
        self.assertEqual(response.context["totals"]["courses"], 1)
        self.assertEqual(response.context["totals"]["enrolled"], 1)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["hp"], 1)

        response = self.client.get(reverse("my_courses"), {"q": "COURSE INSTRUCTOR"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'aria-label="Filter course job cards"')
        self.assertContains(response, "Course name")
        self.assertContains(response, "Date(s)")
        self.assertContains(response, "Enrolled <strong>1</strong>")
        self.assertContains(response, "Passed <strong>1</strong>")

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="course-stage-flow"')
        self.assertContains(response, "Generate stamped list")
        self.assertContains(response, "Upload Stamped List")
        self.assertContains(response, "Generate Front-End")
        self.assertContains(response, "Front-End upload")

    def test_dashboard_combines_history_and_current_without_double_counting(self):
        current_session = self._create_course()
        current_student = Student.objects.create(
            eid="784198012345678",
            name_english="CURRENT STUDENT",
        )
        current_registration = self._registration(current_session)
        current_registration.student = current_student
        current_registration.selected_for_roster = True
        current_registration.assessment_status = "passed"
        current_registration.is_hp = True
        current_registration.save()
        TrainingRecord.objects.create(
            student=current_student,
            session=current_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
            raw_payload={
                "registration_public_id": str(current_registration.public_id)
            },
        )
        pending_registration = self._registration(
            current_session,
            eid="784198112345678",
            name_english="PENDING STUDENT",
            email="pending@example.com",
        )
        pending_registration.selected_for_roster = True
        pending_registration.save(update_fields=["selected_for_roster", "updated_at"])

        historical_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date="2026-02-01",
            end_date="2026-02-02",
            status=CourseSession.Status.COMPLETED,
        )
        CourseInstructor.objects.create(
            session=historical_session,
            instructor=self.instructor,
        )
        historical_student = Student.objects.create(
            eid="784198212345678",
            name_english="HISTORICAL STUDENT",
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.FAIL,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_ttt=True,
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
            duplicate_flag=True,
        )
        Registration.objects.create(
            requested_session=historical_session,
            submitted_name_english="REJECTED STUDENT",
            status=Registration.Status.REJECTED,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["courses"], 2)
        self.assertEqual(response.context["totals"]["enrolled"], 3)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["failed"], 1)
        self.assertEqual(response.context["totals"]["rejected"], 1)
        self.assertEqual(response.context["totals"]["hp"], 1)
        self.assertEqual(response.context["totals"]["ttt"], 1)
        february = response.context["chart_rows"][1]
        september = response.context["chart_rows"][8]
        self.assertEqual(february["enrolled"], 1)
        self.assertEqual(february["failed"], 1)
        self.assertEqual(february["rejected"], 1)
        self.assertEqual(september["enrolled"], 2)
        self.assertEqual(september["passed"], 1)

    def test_dashboard_keeps_date_range_and_offers_all_history(self):
        current_session = self._create_course()
        current_student = Student.objects.create(
            eid="784198312345678",
            name_english="CURRENT YEAR STUDENT",
        )
        TrainingRecord.objects.create(
            student=current_student,
            session=current_session,
            result=TrainingRecord.Result.FAIL,
        )
        historical_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date="2024-03-10",
            end_date="2024-03-11",
            status=CourseSession.Status.COMPLETED,
        )
        historical_student = Student.objects.create(
            eid="784198412345678",
            name_english="OLDER HISTORY STUDENT",
        )
        TrainingRecord.objects.create(
            student=historical_student,
            session=historical_session,
            result=TrainingRecord.Result.PASS,
        )

        self.client.force_login(self.admin)
        current_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertFalse(current_response.context["all_history"])
        self.assertEqual(current_response.context["totals"]["courses"], 1)
        self.assertEqual(current_response.context["totals"]["enrolled"], 1)
        self.assertEqual(current_response.context["totals"]["failed"], 1)

        all_response = self.client.get(
            reverse("dashboard"),
            {
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
                "all_history": "1",
            },
        )
        self.assertTrue(all_response.context["all_history"])
        self.assertContains(all_response, 'name="all_history" value="1" checked')
        self.assertEqual(all_response.context["totals"]["courses"], 2)
        self.assertEqual(all_response.context["totals"]["enrolled"], 2)
        self.assertEqual(all_response.context["totals"]["passed"], 1)
        self.assertEqual(all_response.context["totals"]["failed"], 1)

        historical_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2024-01-01", "date_to": "2024-12-31"},
        )
        self.assertFalse(historical_response.context["all_history"])
        self.assertEqual(historical_response.context["totals"]["courses"], 1)
        self.assertEqual(historical_response.context["totals"]["enrolled"], 1)
        self.assertEqual(historical_response.context["totals"]["passed"], 1)
        self.assertEqual(historical_response.context["totals"]["failed"], 0)

    def test_dashboard_includes_deduplicated_historical_evidence(self):
        session = self._create_course()
        current_student = Student.objects.create(
            eid="784198012345678",
            name_english="CURRENT STUDENT",
        )
        TrainingRecord.objects.create(
            student=current_student,
            session=session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
            is_hp=True,
        )

        source_file = self._history_source_file()
        historical_registration = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198112345678",
            raw_payload={
                "date_التاريخ": {
                    "value": "2026-08-03T00:00:00",
                    "excel_type": "datetime",
                },
                "camp_المعسكر": "Historical Camp",
                "unit_location_الوحدة_الموقع": "PG",
                "instructor": "Course Instructor",
            },
            import_status=SourceRecord.ImportStatus.CONVERTED,
        )
        self._history_proposal(
            historical_registration,
            start_date="2026-08-03",
            course=self.course,
        )

        attendance = self._history_record(
            source_file,
            record_type="course_attendance_evidence",
            eid="784198212345678",
            raw_payload={"session1": "YES", "session2": "YES"},
        )
        result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198212345678",
            raw_payload={"comment": "PASS"},
        )
        self._history_proposal(
            attendance,
            result,
            start_date="2026-08-19",
            course=self.course,
        )

        self._history_record(
            source_file,
            record_type="ttt_candidate_evidence",
            eid="784198312345678",
            raw_payload={
                "training_date": "16-Jul_26",
                "camp_name": "Historical Camp",
                "unit": "PG",
                "instructor_1": "Course Instructor",
            },
        )
        ignored_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198412345678",
            raw_payload={"comment": "FAIL"},
            import_status=SourceRecord.ImportStatus.IGNORED,
        )
        self._history_proposal(
            ignored_result,
            start_date="2026-08-20",
            course=self.course,
        )
        Student.objects.create(
            eid="784198512345678",
            name_english="PROFILE WITHOUT PARTICIPATION",
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )

        totals = response.context["totals"]
        self.assertEqual(totals["courses"], 1)
        self.assertEqual(totals["enrolled"], 3)
        self.assertEqual(totals["passed"], 2)
        self.assertEqual(totals["failed"], 0)
        self.assertEqual(totals["hp"], 1)
        self.assertEqual(totals["ttt"], 1)
        self.assertEqual(response.context["chart_rows"][7]["enrolled"], 2)
        self.assertEqual(response.context["chart_rows"][7]["passed"], 1)

        course_response = self.client.get(
            reverse("dashboard"),
            {
                "course": str(self.course.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(course_response.context["totals"]["enrolled"], 3)

        other_course = Course.objects.create(
            code="OTHER",
            title_english="Other Course",
        )
        other_response = self.client.get(
            reverse("dashboard"),
            {
                "course": str(other_course.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(other_response.context["totals"]["enrolled"], 0)

    def test_dashboard_historical_evidence_respects_dates_and_all_history(self):
        source_file = self._history_source_file()
        current_year = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198612345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        self._history_proposal(
            current_year,
            start_date="2026-08-03",
            course=self.course,
        )
        old_attendance = self._history_record(
            source_file,
            record_type="course_attendance_evidence",
            eid="784198712345678",
            raw_payload={"session1": "YES", "session2": "YES"},
        )
        old_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid="784198712345678",
            raw_payload={"comment": "PASS"},
        )
        self._history_proposal(
            old_attendance,
            old_result,
            start_date="2024-03-10",
            course=self.course,
        )
        self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784198812345678",
            raw_payload={"camp_المعسكر": "Undated Camp"},
        )

        self.client.force_login(self.admin)
        current_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(current_response.context["totals"]["enrolled"], 1)
        self.assertEqual(current_response.context["totals"]["passed"], 0)

        all_response = self.client.get(
            reverse("dashboard"),
            {
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
                "all_history": "1",
            },
        )
        self.assertEqual(all_response.context["totals"]["enrolled"], 3)
        self.assertEqual(all_response.context["totals"]["passed"], 1)
        self.assertEqual(all_response.context["chart_rows"][2]["enrolled"], 1)
        self.assertEqual(all_response.context["chart_rows"][2]["passed"], 1)

        old_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2024-01-01", "date_to": "2024-12-31"},
        )
        self.assertEqual(old_response.context["totals"]["enrolled"], 1)
        self.assertEqual(old_response.context["totals"]["passed"], 1)

    def test_dashboard_canonical_training_precedes_source_evidence(self):
        session = self._create_course()
        student = Student.objects.create(
            eid="784198912345678",
            name_english="CANONICAL STUDENT",
        )
        training = TrainingRecord.objects.create(
            student=student,
            session=session,
            attendance=TrainingRecord.Attendance.PRESENT,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.COMPLETED,
        )
        source_file = self._history_source_file()
        source_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid=student.eid,
            raw_payload={"comment": "FAIL"},
            linked_training_record=training,
        )
        proposal = self._history_proposal(
            source_result,
            start_date="2026-09-01",
            course=self.course,
        )
        proposal.approved_session = session
        proposal.save(update_fields=["approved_session", "updated_at"])

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 1)
        self.assertEqual(response.context["totals"]["passed"], 1)
        self.assertEqual(response.context["totals"]["failed"], 0)

    def test_dashboard_linked_rejection_is_not_counted_twice(self):
        session = self._create_course()
        source_file = self._history_source_file()
        rejected_registration = Registration.objects.create(
            requested_session=session,
            submitted_name_english="REJECTED STUDENT",
            eid_raw="784-1994-1234567-8",
            eid_normalized="784199412345678",
            status=Registration.Status.REJECTED,
            selected_for_roster=True,
            source_file=source_file,
        )
        source = self._history_record(
            source_file,
            record_type="registration_submission",
            eid=rejected_registration.eid_normalized,
            raw_payload={"date_\u0627\u0644\u062a\u0627\u0631\u064a\u062e": "2026-09-01"},
            linked_registration=rejected_registration,
        )
        self._history_proposal(
            source,
            start_date="2026-09-01",
            course=self.course,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 0)
        self.assertEqual(response.context["totals"]["rejected"], 1)

    def test_dashboard_history_survives_approval_and_pending_duplicate_review(self):
        session = self._create_course()
        source_file = self._history_source_file()
        first_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199312345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        repeated_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199312345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        self._history_proposal(
            first_submission,
            repeated_submission,
            start_date="2026-08-03",
            course=self.course,
        )

        approved_submission = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199412345678",
            raw_payload={"date_التاريخ": "2026-09-01"},
        )
        approved_proposal = self._history_proposal(
            approved_submission,
            start_date="2026-09-01",
            course=self.course,
        )
        approved_proposal.approved_session = session
        approved_proposal.proposal_status = (
            CourseSessionProposal.ProposalStatus.APPROVED
        )
        approved_proposal.save(
            update_fields=[
                "approved_session",
                "proposal_status",
                "updated_at",
            ]
        )

        cancelled_student = Student.objects.create(
            eid="784199512345678",
            name_english="CANCELLED CANONICAL ROW",
        )
        cancelled_training = TrainingRecord.objects.create(
            student=cancelled_student,
            session=session,
            result=TrainingRecord.Result.PASS,
            record_status=TrainingRecord.RecordStatus.CANCELLED,
        )
        retained_result = self._history_record(
            source_file,
            record_type="course_result_evidence",
            eid=cancelled_student.eid,
            raw_payload={"comment": "PASS"},
            linked_training_record=cancelled_training,
        )
        self._history_proposal(
            retained_result,
            start_date="2026-09-01",
            course=self.course,
        )

        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(response.context["totals"]["enrolled"], 4)
        self.assertEqual(response.context["totals"]["passed"], 1)

    def test_dashboard_historical_team_and_instructor_scope_are_trusted(self):
        team_a = Team.objects.create(name="Historical Team A")
        team_b = Team.objects.create(name="Historical Team B")
        self.unassigned_instructor.team = team_a
        self.unassigned_instructor.save(update_fields=["team", "updated_at"])
        self.instructor.team = team_b
        self.instructor.save(update_fields=["team", "updated_at"])

        source_file = self._history_source_file()
        own = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199012345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
            linked_instructor=self.unassigned_instructor,
        )
        other = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199112345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
            linked_instructor=self.instructor,
        )
        unresolved = self._history_record(
            source_file,
            record_type="registration_submission",
            eid="784199212345678",
            raw_payload={"date_التاريخ": "2026-08-03"},
        )
        for record in (own, other, unresolved):
            self._history_proposal(
                record,
                start_date="2026-08-03",
                course=self.course,
            )

        self.client.force_login(self.admin)
        team_response = self.client.get(
            reverse("dashboard"),
            {
                "team": str(team_a.pk),
                "date_from": "2026-01-01",
                "date_to": "2026-12-31",
            },
        )
        self.assertEqual(team_response.context["totals"]["enrolled"], 1)

        all_teams_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(all_teams_response.context["totals"]["enrolled"], 3)

        # Regular instructors no longer receive Dashboard by default. An
        # explicit View grant still keeps their historical totals scoped.
        StaffPermissionOverride.objects.create(
            user=self.unassigned_user, feature="dashboard", level="view",
        )
        self.client.force_login(self.unassigned_user)
        instructor_response = self.client.get(
            reverse("dashboard"),
            {"date_from": "2026-01-01", "date_to": "2026-12-31"},
        )
        self.assertEqual(instructor_response.context["totals"]["enrolled"], 1)

    def test_existing_student_and_previous_same_course_require_review(self):
        session = self._create_course()
        master = Student.objects.create(
            eid="784198012345678",
            name_english="TEST STUDENT",
            name_arabic="طالب تجريبي",
            email="student@example.com",
        )
        earlier_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            reference_code="ASM-PREVIOUS",
            start_date="2026-08-01",
            end_date="2026-08-02",
            status=CourseSession.Status.COMPLETED,
        )
        TrainingRecord.objects.create(
            student=master,
            session=earlier_session,
            result=TrainingRecord.Result.PASS,
        )
        registration = self._registration(session)
        self.client.force_login(self.instructor_user)

        response = self.client.get(
            reverse("instructor_course_workspace", args=[session.public_id])
        )
        self.assertContains(response, "Existing master student")
        self.assertContains(response, "Previously attended this course")
        self.assertContains(response, "Review required")

        response = self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "pending",
            },
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertFalse(registration.selected_for_roster)

        response = self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"duplicate_decision_{registration.public_id}": "reviewed_allow",
                f"duplicate_notes_{registration.public_id}": "Identity checked",
                f"result_{registration.public_id}": "pending",
            },
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertTrue(registration.selected_for_roster)
        self.assertEqual(
            registration.duplicate_review_status,
            Registration.DuplicateReviewStatus.ALLOWED,
        )
        self.assertEqual(registration.duplicate_review_notes, "Identity checked")
        self.assertEqual(len(registration.duplicate_review_fingerprint), 64)
        self.assertEqual(registration.duplicate_reviewed_by, self.instructor_user)

        self.client.force_login(self.admin)
        response = self.client.get(reverse("data_quality") + "?state=allowed")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Identity checked")
        self.assertContains(response, "Previously attended this course")

    def test_external_upload_contains_only_passed_selected_student(self):
        session = self._create_course()
        registration = self._registration(session)
        self.client.force_login(self.instructor_user)
        self.client.post(
            reverse("course_roster_save", args=[session.public_id]),
            {
                f"selected_{registration.public_id}": "yes",
                f"result_{registration.public_id}": "passed",
            },
        )
        response = self.client.get(
            reverse("external_upload_xlsx", args=[session.public_id])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            "attachment; filename*=UTF-8''"
            "01-02_September_COURSE_Test_Camp_Students%20upload.xlsx",
        )
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["obs", "TCCC ASM"])
        sheet = workbook["TCCC ASM"]
        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 6)],
            ["email", "emirates_id", "full_name", "session1", "session2"],
        )
        self.assertEqual(sheet["C2"].value, "طالب تجريبي")
        self.assertEqual(sheet["B2"].value, "784-1980-1234567-8")
        self.assertEqual(sheet["D2"].value, "YES")
        self.assertEqual(sheet["E2"].value, "YES")

    def test_day_one_filenames_match_supplied_examples(self):
        session = self._create_course()
        self.instructor.name_english = "Vasilios Example"
        self.instructor.save(update_fields=["name_english", "updated_at"])
        self.unassigned_instructor.name_english = "Abdulsattar Dalgamuni"
        self.unassigned_instructor.save(
            update_fields=["name_english", "updated_at"]
        )
        CourseInstructor.objects.create(
            session=session,
            instructor=self.unassigned_instructor,
        )
        session.start_date = timezone.datetime(2026, 8, 19).date()
        session.end_date = timezone.datetime(2026, 8, 20).date()
        self.camp.name = "Specialized Reserve Nahel PG"
        self.camp.save(update_fields=["name", "updated_at"])

        expected_base = (
            "19-20_August_Vasilios-Abdulsattar_"
            "Specialized_Reserve_Nahel_PG"
        )
        self.assertEqual(day_one_list_filename(session), expected_base + ".xlsx")
        self.assertEqual(
            day_one_students_upload_filename(session),
            expected_base + "_Students upload.xlsx",
        )

    def test_course_lists_use_compact_human_references(self):
        session = self._create_course()
        self.client.force_login(self.admin)

        response = self.client.get(reverse("my_courses"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="job-card job-status-')
        self.assertContains(response, "View Course")
        self.assertContains(response, "Location")
        self.assertContains(response, "Test Camp")
        self.assertContains(response, 'class="job-instructors"')
        self.assertContains(response, "Instructor(s)")
        self.assertContains(response, "COURSE INSTRUCTOR")
        self.assertNotContains(response, "ASM · TCCC · 01 SEP 2026")
        self.assertNotContains(response, session.reference_code)

        response = self.client.get(reverse("courses_sessions"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Master Course Directory")
        self.assertContains(response, "ASM - TCCC")
        self.assertContains(response, "Course Directory")
        self.assertContains(response, 'class="course-directory-row"')
        self.assertContains(response, 'class="course-directory-date-cell"', count=2)
        self.assertContains(response, 'class="directory-instructor-list"')
        self.assertContains(response, 'class="directory-badge lifecycle-')
        self.assertContains(
            response,
            reverse("instructor_course_workspace", args=[session.public_id]),
        )

    def test_reports_offer_focused_one_row_views(self):
        session = self._create_course()
        registration = self._registration(session)
        registration.selected_for_roster = True
        registration.assessment_status = "passed"
        registration.is_hp = True
        registration.save(
            update_fields=[
                "selected_for_roster",
                "assessment_status",
                "is_hp",
                "updated_at",
            ]
        )
        item = InventoryItem.objects.create(
            name="Report Kit",
            category=InventoryItem.Category.EQUIPMENT,
            unit="kit",
        )
        CourseInstructorInventoryUsage.objects.create(
            session=session,
            instructor=self.instructor,
            item=item,
            quantity_used=2,
        )
        self.client.force_login(self.admin)

        for view_name, expected in [
            ("progress", "Course progress"),
            ("outcomes", "Student outcomes"),
            ("documents", "Document status"),
            ("instructors", "COURSE INSTRUCTOR"),
            ("inventory", "Report Kit"),
            ("activity", "Activity log"),
        ]:
            response = self.client.get(
                reverse("reports") + f"?view={view_name}"
            )
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, expected)

    def test_student_directory_print_uses_day_one_browser_sheet(self):
        student = Student.objects.create(
            eid="784198012345678",
            name_english="PRINT TEST STUDENT",
            name_arabic="طالب للطباعة",
            email="print@example.com",
        )
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("students_print"),
            {
                "mode": "selected",
                "student_ids": [student.pk],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Student Directory List")
        self.assertContains(response, "EID / رقم الهوية")
        self.assertContains(response, "Day 1")
        self.assertContains(response, "Day 2")
        self.assertContains(response, "Status")
        self.assertContains(response, "PRINT TEST STUDENT")
        self.assertTrue(
            ActivityLog.objects.filter(
                action=ActivityLog.Action.PRINT,
                object_type="Student",
            ).exists()
        )

    def test_admin_navigation_exposes_course_directory(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'href="' + reverse("courses_sessions") + '"')
        self.assertContains(response, "iqarus-logo-approved")

        self.client.force_login(self.unassigned_user)
        response = self.client.get(reverse("courses_sessions"))
        self.assertEqual(response.status_code, 403)


class SecurityAndPlannerTests(TestCase):
    def test_email_addresses_are_normalized_at_model_write_boundaries(self):
        User = get_user_model()
        user = User.objects.create_user(
            username="staff-user",
            email="STAFF.USER@EXAMPLE.COM",
            password="safe-test-password",
        )
        instructor = Instructor.objects.create(
            user=user,
            name_english="STAFF USER",
            email="INSTRUCTOR@EXAMPLE.COM",
        )
        student = Student.objects.create(
            eid="784198012345678",
            name_english="EMAIL TEST",
            email="STUDENT@EXAMPLE.COM",
        )
        registration = Registration.objects.create(
            submitted_name_english="EMAIL TEST",
            email_raw="REGISTRATION@EXAMPLE.COM",
        )

        user.refresh_from_db()
        instructor.refresh_from_db()
        student.refresh_from_db()
        registration.refresh_from_db()
        self.assertEqual(user.email, "staff.user@example.com")
        self.assertEqual(instructor.email, "instructor@example.com")
        self.assertEqual(student.email, "student@example.com")
        self.assertEqual(registration.email_raw, "registration@example.com")

    def test_temporary_password_account_is_limited_to_password_change(self):
        User = get_user_model()
        user = User.objects.create_user(
            username="temporary-user",
            password="safe-test-password",
        )
        UserSecurityProfile.objects.create(
            user=user,
            must_change_password=True,
        )
        self.client.force_login(user)

        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, reverse("password_change"))
        response = self.client.get(reverse("password_change"))
        self.assertEqual(response.status_code, 200)

    def test_j35_planner_group_can_create_an_allocation(self):
        User = get_user_model()
        manager = User.objects.create_user(
            username="j35-manager",
            password="safe-test-password",
        )
        manager.groups.add(Group.objects.create(name=J35_PLANNER_GROUP))
        staff_user = User.objects.create_user(
            username="allocated-instructor",
            password="safe-test-password",
        )
        instructor = Instructor.objects.create(
            user=staff_user,
            name_english="ALLOCATED INSTRUCTOR",
            email="allocated@example.com",
        )
        self.client.force_login(manager)

        response = self.client.post(
            reverse("j35_planner"),
            {
                "action": "create",
                "instructor": instructor.id,
                "session": "",
                "camp": "",
                "activity": "J35 COURSE SUPPORT",
                "start_date": "2026-09-01",
                "end_date": "2026-09-03",
                "status": InstructorAllocation.Status.PUBLISHED,
                "notes": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            InstructorAllocation.objects.filter(
                instructor=instructor,
                activity="J35 COURSE SUPPORT",
                status=InstructorAllocation.Status.PUBLISHED,
            ).exists()
        )

    def test_j35_course_pipeline_creates_all_selected_instructors(self):
        User = get_user_model()
        manager = User.objects.create_user(
            username="pipeline-manager",
            password="safe-test-password",
        )
        manager.groups.add(Group.objects.create(name=J35_PLANNER_GROUP))
        camp = Camp.objects.create(name="Pipeline Camp")
        course = Course.objects.create(code="PIPE", title_english="Pipeline Course")
        session = CourseSession.objects.create(
            course=course,
            camp=camp,
            start_date=date(2026, 9, 8),
            end_date=date(2026, 9, 9),
        )
        first_user = User.objects.create_user(
            username="pipeline-one",
            password="safe-test-password",
        )
        second_user = User.objects.create_user(
            username="pipeline-two",
            password="safe-test-password",
        )
        first = Instructor.objects.create(
            user=first_user,
            name_english="PIPELINE ONE",
            email="one@example.com",
        )
        second = Instructor.objects.create(
            user=second_user,
            name_english="PIPELINE TWO",
            email="two@example.com",
        )
        self.client.force_login(manager)

        response = self.client.post(
            reverse("j35_planner"),
            {
                "action": "pipeline_create",
                "session_id": session.id,
                "instructor_ids": [first.id, second.id],
                "status": InstructorAllocation.Status.PUBLISHED,
                "activity": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 302)
        allocations = InstructorAllocation.objects.filter(session=session)
        self.assertEqual(allocations.count(), 2)
        self.assertEqual(
            set(allocations.values_list("instructor_id", flat=True)),
            {first.id, second.id},
        )
        self.assertTrue(
            ActivityLog.objects.filter(
                object_type="InstructorAllocation",
                details__source="course_pipeline",
            ).exists()
        )

    def test_staff_account_can_keep_its_initial_password(self):
        User = get_user_model()
        admin = User.objects.create_user(
            username="staff-admin",
            password="safe-test-password",
            is_staff=True,
            is_superuser=True,
        )
        self.client.force_login(admin)

        response = self.client.post(
            reverse("staff"),
            {
                "action": "add_instructor",
                "name": "OPTIONAL PASSWORD INSTRUCTOR",
                "email": "OPTIONAL.PASSWORD@EXAMPLE.COM",
                "username": "optional-password-instructor",
                "temporary_password": "Dune^River9!Quartz2026",
            },
        )

        self.assertEqual(response.status_code, 302)
        instructor = Instructor.objects.get(
            name_english="OPTIONAL PASSWORD INSTRUCTOR"
        )
        profile = UserSecurityProfile.objects.get(user=instructor.user)
        self.assertFalse(profile.must_change_password)
        self.assertIsNone(profile.temporary_password_issued_at)

    def test_operations_admin_can_manage_j35_under_current_staff_rules(self):
        """The approved layout grants planning to administrators and planners."""

        User = get_user_model()
        operations_user = User.objects.create_user(
            username="operations-user",
            password="safe-test-password",
        )
        operations_user.groups.add(
            Group.objects.create(name=OPERATIONS_ADMIN_GROUP)
        )
        self.assertTrue(can_manage_j35_planner(operations_user))

        operations_user.groups.add(
            Group.objects.create(name=J35_PLANNER_GROUP)
        )
        self.assertTrue(can_manage_j35_planner(operations_user))

    def test_lead_instructor_group_sync_grants_narrow_j35_planner_role(self):
        User = get_user_model()
        user = User.objects.create_user(
            username="lead-instructor",
            password="safe-test-password",
        )
        instructor = Instructor.objects.create(
            user=user,
            name_english="LEAD INSTRUCTOR",
            email="lead@example.com",
        )
        instructor.roles.add(InstructorRole.objects.create(name="LI"))

        synchronise_staff_groups(instructor, ensure_predefined_groups())

        self.assertTrue(
            user.groups.filter(name=J35_PLANNER_GROUP).exists()
        )
        self.assertTrue(can_manage_j35_planner(user))

    def test_j35_group_has_only_needed_planner_model_permissions(self):
        group = ensure_predefined_groups()[J35_PLANNER_GROUP]
        permissions = set(group.permissions.values_list("codename", flat=True))

        for model in (CampContact, J35TrainingNeed, InstructorAllocation):
            for action in ("add", "change", "view"):
                self.assertIn(f"{action}_{model._meta.model_name}", permissions)
            self.assertNotIn(f"delete_{model._meta.model_name}", permissions)

    def test_staff_group_sync_command_is_safe_by_default_and_idempotent(self):
        User = get_user_model()
        user = User.objects.create_user(
            username="command-lead",
            password="safe-test-password",
        )
        instructor = Instructor.objects.create(
            user=user,
            name_english="COMMAND LEAD",
            email="command-lead@example.com",
        )
        instructor.roles.add(
            InstructorRole.objects.create(name="Lead Instructor")
        )

        output = StringIO()
        call_command("synchronise_staff_groups", stdout=output)
        self.assertIn("Dry run only", output.getvalue())
        self.assertFalse(user.groups.filter(name=J35_PLANNER_GROUP).exists())

        call_command("synchronise_staff_groups", "--apply", stdout=StringIO())
        self.assertTrue(user.groups.filter(name=J35_PLANNER_GROUP).exists())

        repeat_output = StringIO()
        call_command("synchronise_staff_groups", "--apply", stdout=repeat_output)
        self.assertIn("synchronised 0", repeat_output.getvalue())


class NotificationInboxTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(
            username="notification-owner",
            password="safe-test-password",
        )
        self.other_user = User.objects.create_user(
            username="notification-other",
            password="safe-test-password",
        )

    def test_inbox_is_private_and_header_shows_only_own_unread_count(self):
        Notification.objects.create(
            recipient=self.user,
            title="Your J35 allocation",
            message="You have a new allocation.",
            link="/my-courses/",
        )
        Notification.objects.create(
            recipient=self.other_user,
            title="Other user's private message",
            message="This must not be visible.",
            link="https://example.invalid/outside",
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("notification_inbox"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Your J35 allocation")
        self.assertNotContains(response, "Other user's private message")
        self.assertEqual(response.context["unread_count"], 1)
        self.assertEqual(response.context["notification_unread_count"], 1)
        self.assertContains(response, 'class="notification-badge"')

    def test_mark_read_is_post_only_and_scoped_to_current_user(self):
        own = Notification.objects.create(
            recipient=self.user,
            title="Own notification",
            message="Read this.",
        )
        other = Notification.objects.create(
            recipient=self.other_user,
            title="Other notification",
            message="Do not read this.",
        )
        self.client.force_login(self.user)

        get_response = self.client.get(
            reverse("notification_mark_read", args=[own.id])
        )
        self.assertEqual(get_response.status_code, 405)

        forbidden_response = self.client.post(
            reverse("notification_mark_read", args=[other.id])
        )
        self.assertEqual(forbidden_response.status_code, 404)
        other.refresh_from_db()
        self.assertIsNone(other.read_at)

        response = self.client.post(
            reverse("notification_mark_read", args=[own.id]),
            {"next": "https://example.invalid/not-allowed"},
        )
        self.assertRedirects(response, reverse("notification_inbox"))
        own.refresh_from_db()
        self.assertIsNotNone(own.read_at)

    def test_inbox_does_not_render_external_notification_links(self):
        Notification.objects.create(
            recipient=self.user,
            title="Unsafe legacy link",
            message="The link must be suppressed.",
            link="https://example.invalid/outside",
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("notification_inbox"))

        self.assertNotContains(response, 'href="https://example.invalid/outside"')


class J35WeeklyPipelineTests(TestCase):
    """Regression coverage for the durable weekly planning workflow."""

    def setUp(self):
        User = get_user_model()
        self.manager = User.objects.create_user(
            username="weekly-j35-manager",
            password="safe-test-password",
        )
        self.manager.groups.add(Group.objects.create(name=J35_PLANNER_GROUP))
        self.camp = Camp.objects.create(name="Weekly Pipeline Camp", area="North")
        self.course = Course.objects.create(
            code="WEEKLY",
            title_english="Weekly Pipeline Course",
        )
        self.first_user = User.objects.create_user(
            username="weekly-instructor-one",
            password="safe-test-password",
        )
        self.second_user = User.objects.create_user(
            username="weekly-instructor-two",
            password="safe-test-password",
        )
        self.first = Instructor.objects.create(
            user=self.first_user,
            name_english="WEEKLY INSTRUCTOR ONE",
            email="weekly-one@example.com",
        )
        self.second = Instructor.objects.create(
            user=self.second_user,
            name_english="WEEKLY INSTRUCTOR TWO",
            email="weekly-two@example.com",
        )
        self.client.force_login(self.manager)

    def _need(self, **overrides):
        values = {
            "camp": self.camp,
            "course": self.course,
            "source": J35TrainingNeed.Source.CLIENT_REQUEST,
            "status": J35TrainingNeed.Status.REQUESTED,
            "priority": J35TrainingNeed.Priority.HIGH,
            "requested_start_date": date(2026, 9, 21),
            "requested_end_date": date(2026, 9, 22),
            "required_instructors": 1,
            "created_by": self.manager,
            "updated_by": self.manager,
        }
        values.update(overrides)
        return J35TrainingNeed.objects.create(**values)

    def test_compact_need_post_forces_safe_lifecycle_fields(self):
        draft_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date=date(2026, 10, 1),
            end_date=date(2026, 10, 2),
        )

        response = self.client.post(
            reverse("j35_planner"),
            {
                "action": "create_need",
                "camp_id": self.camp.id,
                "course_id": self.course.id,
                "request_type": "client_request",
                "priority": "urgent",
                "start_date": "2026-09-21",
                "end_date": "2026-09-22",
                "required_instructors": "2",
                # These are deliberately forged lifecycle controls. The compact
                # planner may not use them to bypass confirmation.
                "status": J35TrainingNeed.Status.CONFIRMED,
                "proposed_session": draft_session.id,
            },
        )

        self.assertEqual(response.status_code, 302)
        need = J35TrainingNeed.objects.get(camp=self.camp)
        self.assertEqual(need.status, J35TrainingNeed.Status.REQUESTED)
        self.assertIsNone(need.proposed_session_id)
        self.assertEqual(need.requested_start_date, date(2026, 9, 21))
        self.assertEqual(need.requested_end_date, date(2026, 9, 22))

    def test_need_update_cannot_overwrite_server_owned_lifecycle_fields(self):
        linked_draft = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date=date(2026, 10, 1),
            end_date=date(2026, 10, 2),
        )
        need = self._need(
            status=J35TrainingNeed.Status.PROPOSED,
            proposed_session=linked_draft,
            contacted_at=timezone.now(),
        )
        original_contacted_at = need.contacted_at

        response = self.client.post(
            reverse("j35_planner"),
            {
                "action": "update_need",
                "need_id": str(need.public_id),
                "camp_id": self.camp.id,
                "course_id": self.course.id,
                "request_type": "outreach",
                "priority": "normal",
                "start_date": "2026-09-24",
                "end_date": "2026-09-25",
                "required_instructors": "1",
                # These fields must never be client-controlled through the
                # editable planner form.
                "status": J35TrainingNeed.Status.CONFIRMED,
                "proposed_session": "",
                "contacted_at": "2001-01-01T00:00",
            },
        )

        self.assertEqual(response.status_code, 302)
        need.refresh_from_db()
        self.assertEqual(need.status, J35TrainingNeed.Status.PROPOSED)
        self.assertEqual(need.proposed_session_id, linked_draft.id)
        self.assertEqual(need.contacted_at, original_contacted_at)
        self.assertEqual(need.requested_start_date, date(2026, 9, 24))

    def test_confirmation_creates_canonical_assignments_notifications_and_audit(self):
        need = self._need(required_instructors=2)

        response = self.client.post(
            reverse("j35_planner"),
            {
                "action": "confirm_need",
                "need_id": str(need.public_id),
                "instructor_ids": [self.first.id, self.second.id],
            },
        )

        self.assertEqual(response.status_code, 302)
        need.refresh_from_db()
        self.assertEqual(need.status, J35TrainingNeed.Status.CONFIRMED)
        self.assertIsNotNone(need.proposed_session_id)
        assignments = CourseInstructor.objects.filter(session=need.proposed_session)
        self.assertEqual(assignments.count(), 2)
        self.assertTrue(all(assignment.notified_at for assignment in assignments))
        allocations = InstructorAllocation.objects.filter(training_need=need)
        self.assertEqual(allocations.count(), 2)
        self.assertTrue(
            all(
                allocation.allocation_kind
                == InstructorAllocation.AllocationKind.COURSE
                for allocation in allocations
            )
        )
        self.assertEqual(
            Notification.objects.filter(
                recipient__in=[self.first_user, self.second_user],
                title="J35 training allocation confirmed",
            ).count(),
            2,
        )
        self.assertTrue(
            ActivityLog.objects.filter(
                action=ActivityLog.Action.APPROVE,
                object_type="J35TrainingNeed",
                object_id=str(need.public_id),
            ).exists()
        )

    def test_conflict_requires_audited_override_and_confirmed_plan_cannot_cancel(self):
        InstructorAllocation.objects.create(
            instructor=self.first,
            activity="Existing duty",
            start_date=date(2026, 9, 21),
            end_date=date(2026, 9, 22),
            status=InstructorAllocation.Status.PUBLISHED,
            created_by=self.manager,
            updated_by=self.manager,
        )
        need = self._need()

        blocked = self.client.post(
            reverse("j35_planner"),
            {
                "action": "confirm_need",
                "need_id": str(need.public_id),
                "instructor_ids": [self.first.id],
            },
        )
        self.assertEqual(blocked.status_code, 302)
        need.refresh_from_db()
        self.assertNotEqual(need.status, J35TrainingNeed.Status.CONFIRMED)

        confirmed = self.client.post(
            reverse("j35_planner"),
            {
                "action": "confirm_need",
                "need_id": str(need.public_id),
                "instructor_ids": [self.first.id],
                "override_reason": "Client specifically requested this trainer.",
            },
        )
        self.assertEqual(confirmed.status_code, 302)
        allocation = InstructorAllocation.objects.get(training_need=need, instructor=self.first)
        self.assertEqual(
            allocation.override_reason,
            "Client specifically requested this trainer.",
        )

        cancelled = self.client.post(
            reverse("j35_planner"),
            {
                "action": "cancel",
                "allocation_id": str(allocation.public_id),
            },
        )
        self.assertEqual(cancelled.status_code, 302)
        allocation.refresh_from_db()
        self.assertEqual(allocation.status, InstructorAllocation.Status.PUBLISHED)

    def test_recent_canonical_course_workload_affects_recommendation_order(self):
        historic_session = CourseSession.objects.create(
            course=self.course,
            camp=self.camp,
            start_date=date(2026, 9, 7),
            end_date=date(2026, 9, 8),
            status=CourseSession.Status.REGISTRATION_CLOSED,
        )
        CourseInstructor.objects.create(
            session=historic_session,
            instructor=self.first,
        )
        need = self._need()

        cards = recommendation_cards(
            need,
            week_start=date(2026, 9, 21),
            instructors=[self.first, self.second],
        )

        self.assertEqual(cards[0].instructor, self.second)
        self.assertGreater(cards[0].score, cards[1].score)
