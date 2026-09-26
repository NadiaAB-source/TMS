"""Strict workbook reader. No Django/database writes in this module."""
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePosixPath
import hashlib
import re
import zipfile

from openpyxl import load_workbook

VERSION = 'IQARUS-HISTORY-1'
MAX_UPLOAD = 100 * 1024 * 1024
MAX_EXPANDED = 250 * 1024 * 1024
MAX_DOCUMENT = 25 * 1024 * 1024
MAX_ROWS = 50000
SCHEMA = {
    'Courses': 'course_key course_code course_name start_date end_date camp_name capacity status registration_open service_branch service_branch_other students_per_instructor poc_name poc_phone map_url notes'.split(),
    'Students': 'student_key emirates_id name_english name_arabic email phone identity_status'.split(),
    'Enrolments': 'course_key student_key registration_status selected day1 day2 result hp ttt remarks duplicate_decision duplicate_notes submitted_at unit'.split(),
    'Trainers': 'course_key username role notified_at acknowledged_at conflict_reason j35_color'.split(),
    'CourseHistory': 'course_key created_at created_by confirmed_at information_sent_at frontend_generated_at frontend_confirmed_at frontend_confirmed_by day1_list_created_at'.split(),
    'Materials': 'course_key username item_name quantity_used quantity_consumed quantity_deteriorated stock_treatment notes'.split(),
    'Documents': 'course_key file_name uploaded_at uploaded_by notes'.split(),
}
REQUIRED = {
    'Courses': {'course_key', 'course_code', 'course_name', 'start_date', 'end_date', 'camp_name', 'capacity', 'status'},
    'Students': {'student_key', 'name_english', 'identity_status'},
    'Enrolments': {'course_key', 'student_key', 'registration_status', 'selected', 'day1', 'day2', 'result', 'hp', 'ttt'},
    'Trainers': {'course_key', 'username', 'role'},
    'CourseHistory': {'course_key'},
    'Materials': {'course_key', 'username', 'item_name', 'quantity_used', 'quantity_consumed', 'quantity_deteriorated', 'stock_treatment'},
    'Documents': {'course_key', 'file_name'},
}


class ImportProblem(ValueError):
    pass


def text(value):
    return '' if value is None else str(value).strip()


def fail(row, message):
    raise ImportProblem(f"{row.get('_sheet', 'Workbook')} row {row.get('_row', '?')}: {message}")


def choice(row, field, options, default=''):
    value = text(row.get(field)).lower() or default
    if value not in options:
        fail(row, f'{field} must be one of: ' + ', '.join(options))
    return value


def boolean(row, field, default=False):
    value = row.get(field)
    if value in (None, ''):
        return default
    if isinstance(value, bool):
        return value
    token = text(value).lower()
    if token in ('yes', 'true', '1'):
        return True
    if token in ('no', 'false', '0'):
        return False
    fail(row, f'{field} must be yes or no.')


def integer(row, field, minimum=0, maximum=2147483647, default=None):
    value = row.get(field)
    if value in (None, '') and default is not None:
        return default
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or not minimum <= number <= maximum:
            raise ValueError
        return int(number)
    except (InvalidOperation, TypeError, ValueError):
        fail(row, f'{field} must be a whole number from {minimum} to {maximum}.')


def date_value(row, field):
    value = row.get(field)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(text(value))
    except ValueError:
        fail(row, f'{field} needs an Excel date or YYYY-MM-DD.')


def timestamp(row, field):
    value = row.get(field)
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        # Excel has no timezone field. The template explicitly labels its
        # timestamp convention as UAE local time (UTC+04:00).
        return value if value.tzinfo is not None else value.replace(tzinfo=timezone(timedelta(hours=4)))
    try:
        parsed = datetime.fromisoformat(text(value).replace('Z', '+00:00'))
        if parsed.tzinfo is None:
            raise ValueError
        return parsed
    except ValueError:
        fail(row, f'{field} needs a timezone, for example 2026-09-01T09:00:00+04:00.')


def key(row, field):
    value = text(row.get(field))
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,69}', value):
        fail(row, f'{field} needs a unique short code using letters, numbers, dots, dashes or underscores.')
    return value


def zip_entries(data, *, limit=MAX_EXPANDED):
    try:
        z = zipfile.ZipFile(BytesIO(data))
        entries = z.infolist()
        if len(entries) > 10000 or sum(i.file_size for i in entries) > limit:
            raise ImportProblem('The ZIP expands beyond the allowed size.')
        names = set()
        for item in entries:
            path = PurePosixPath(item.filename)
            if (path.is_absolute() or '..' in path.parts or '\\' in item.filename or
                ':' in item.filename or item.filename in names or item.flag_bits & 1 or
                (item.external_attr >> 16) & 0o170000 == 0o120000):
                raise ImportProblem('ZIP contains an unsafe, encrypted or repeated file path.')
            names.add(item.filename)
        return z
    except zipfile.BadZipFile:
        raise ImportProblem('The upload is not a valid Excel workbook or ZIP.') from None


def read_upload(data, filename):
    if not data or len(data) > MAX_UPLOAD:
        raise ImportProblem('Choose an Excel workbook or ZIP up to 100 MB.')
    documents = {}
    if filename.lower().endswith('.xlsx'):
        workbook = data
    elif filename.lower().endswith('.zip'):
        with zip_entries(data) as z:
            workbooks = [name for name in z.namelist() if name.lower().endswith('.xlsx') and '/' not in name]
            if len(workbooks) != 1:
                raise ImportProblem('Put exactly one .xlsx workbook in the top of the ZIP, with scans in documents/.')
            workbook = z.read(workbooks[0])
            for info in z.infolist():
                if info.is_dir() or info.filename == workbooks[0]:
                    continue
                if not info.filename.startswith('documents/') or len(PurePosixPath(info.filename).parts) != 2:
                    raise ImportProblem('Only the workbook and files directly inside documents/ belong in the import ZIP.')
                if info.file_size > MAX_DOCUMENT:
                    raise ImportProblem('Each stamped-list scan must be 25 MB or smaller.')
                documents[PurePosixPath(info.filename).name] = z.read(info.filename)
    else:
        raise ImportProblem('Choose an .xlsx workbook or a .zip containing the workbook and stamped lists.')
    with zip_entries(workbook):
        pass
    try:
        book = load_workbook(BytesIO(workbook), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ImportProblem('Excel could not be read. Save a normal .xlsx workbook and retry.') from exc
    rows = {name: [] for name in SCHEMA}
    try:
        if any(name not in book.sheetnames for name in ('Courses', 'Students', 'Enrolments', 'Trainers')):
            raise ImportProblem('Keep the Courses, Students, Enrolments and Trainers sheets.')
        if any(name not in set(SCHEMA) | {'ReadMe'} for name in book.sheetnames):
            raise ImportProblem('Unexpected sheet name. Use the supplied template so no data is silently ignored.')
        total = 0
        for name, columns in SCHEMA.items():
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(columns) + 3:
                raise ImportProblem(f'{name}: too many rows or unexpected columns.')
            iterator = sheet.iter_rows()
            first = next(iterator, [])
            headers = [text(c.value) for c in first]
            while headers and headers[-1] == '':
                headers.pop()
            if len(headers) != len(set(headers)) or any(h not in columns for h in headers):
                raise ImportProblem(f'{name}: repeated or unknown column heading. Keep the template headings.')
            if not REQUIRED[name].issubset(headers):
                raise ImportProblem(f'{name}: missing required headings: ' + ', '.join(sorted(REQUIRED[name] - set(headers))))
            for number, cells in enumerate(iterator, 2):
                values = [c.value for c in cells]
                if not any(v not in (None, '') for v in values):
                    continue
                total += 1
                if total > MAX_ROWS:
                    raise ImportProblem('Import at most 50,000 populated rows in one batch.')
                row = dict(zip(headers, values))
                row.update(_sheet=name, _row=number)
                if any(v not in (None, '') for v in values[len(headers):]):
                    fail(row, 'Data exists in a column without a heading.')
                if any(c.data_type == 'f' or (isinstance(c.value, str) and c.value.lstrip().startswith('=')) for c in cells):
                    fail(row, 'Use values, not formulas, in import cells.')
                for required in REQUIRED[name]:
                    if row.get(required) in (None, ''):
                        fail(row, f'{required} is required.')
                rows[name].append(row)
    finally:
        book.close()
    if not rows['Courses']:
        raise ImportProblem('Add at least one course to Courses.')
    return {'rows': rows, 'documents': documents, 'sha256': hashlib.sha256(data).hexdigest(),
            'filename': PurePosixPath(filename.replace('\\', '/')).name, 'size': len(data)}


def serializable(row):
    return {k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in row.items() if not k.startswith('_')}
