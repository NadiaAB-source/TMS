from io import BytesIO
from urllib.parse import quote

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .access import require_session_access
from .course_services import day_one_students_upload_filename
from .models import ActivityLog, CourseSession, Registration
from .student_services import registration_export_values


HEADERS = ["email", "emirates_id", "full_name", "session1", "session2"]


@login_required
def external_upload_xlsx(request, public_id):
    session = get_object_or_404(
        CourseSession.objects.select_related("course", "camp"), public_id=public_id
    )
    require_session_access(request.user, session, edit=True)
    if session.status == CourseSession.Status.CANCELLED:
        raise PermissionDenied("A cancelled course is read-only.")
    registrations = list(
        Registration.objects.filter(
            requested_session=session,
            status=Registration.Status.APPROVED,
            selected_for_roster=True,
            student__isnull=False,
            day1_attended=True,
            day2_attended=True,
            assessment_status="passed",
        )
        .select_related("student")
        .order_by("student__name_english", "id")
    )
    if not registrations:
        return HttpResponse(
            "No selected students are marked Passed.",
            status=400,
            content_type="text/plain; charset=utf-8",
        )
    export_rows = []
    errors = []
    seen_eids = set()
    seen_emails = set()
    for row_number, registration in enumerate(registrations, start=2):
        values = registration_export_values(registration)
        if values["errors"]:
            errors.append(
                f"Row {row_number}: missing or invalid " + ", ".join(values["errors"])
            )
            continue
        if values["emirates_id"] in seen_eids:
            errors.append(f"Row {row_number}: duplicate Emirates ID.")
        if values["email"] in seen_emails:
            errors.append(f"Row {row_number}: duplicate email address.")
        seen_eids.add(values["emirates_id"])
        seen_emails.add(values["email"])
        export_rows.append(
            [values["email"], values["emirates_id"], values["full_name"], "YES", "YES"]
        )
    if errors:
        return HttpResponse(
            "The Excel file was not created because corrections are required:\n"
            + "\n".join(errors),
            status=400,
            content_type="text/plain; charset=utf-8",
        )

    workbook = Workbook()
    workbook.active.title = "obs"
    worksheet = workbook.create_sheet("TCCC ASM")
    workbook.active = 1
    worksheet.sheet_view.zoomScale = 80
    worksheet.sheet_view.zoomScaleNormal = 80
    worksheet.sheet_properties.pageSetUpPr.fitToPage = False
    worksheet.page_setup.orientation = "landscape"
    widths = [43, 17.25, 31.25, 14.375, 15.5]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    black_fill = PatternFill(fill_type="solid", fgColor="000000")
    header_fonts = [
        Font(name="Book Antiqua", size=14),
        Font(name="Book Antiqua", size=14),
        Font(name="Dubai", size=14),
        Font(name="Book Antiqua", size=14),
        Font(name="Book Antiqua", size=14),
    ]
    header_alignments = [
        Alignment(horizontal="left", vertical="center"),
        Alignment(horizontal="center", vertical="center"),
        Alignment(horizontal="left", vertical="center"),
        Alignment(horizontal="left", vertical="center"),
        Alignment(horizontal="left", vertical="center"),
    ]
    worksheet.row_dimensions[1].height = 36.75
    for column, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(1, column, header)
        cell.font = header_fonts[column - 1]
        cell.alignment = header_alignments[column - 1]
        cell.fill = white_fill
        cell.border = border
    data_fonts = [
        Font(name="Calibri", size=14),
        Font(name="Aptos Narrow", size=11),
        Font(name="Aptos Narrow", size=11),
        Font(name="Book Antiqua", size=11),
        Font(name="Book Antiqua", size=11),
    ]
    data_alignments = [
        Alignment(horizontal="left", vertical="center"),
        Alignment(horizontal="center", vertical="center"),
        Alignment(horizontal="right", vertical="center", readingOrder=2),
        Alignment(horizontal="center", vertical="center"),
        Alignment(horizontal="center", vertical="center"),
    ]
    for row_number, values in enumerate(export_rows, start=2):
        worksheet.row_dimensions[row_number].height = 36.75
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = data_fonts[column - 1]
            cell.alignment = data_alignments[column - 1]
            cell.border = border
            cell.number_format = "@"
            cell.fill = black_fill if column in {4, 5} else white_fill
    output = BytesIO()
    workbook.save(output)
    data = output.getvalue()
    verification = load_workbook(BytesIO(data), data_only=False)
    sheet = verification["TCCC ASM"]
    if verification.sheetnames != ["obs", "TCCC ASM"]:
        raise RuntimeError("Workbook sheet structure verification failed.")
    if sheet.max_row != len(export_rows) + 1 or sheet.max_column != 5:
        raise RuntimeError("Workbook size verification failed.")
    if [sheet.cell(1, column).value for column in range(1, 6)] != HEADERS:
        raise RuntimeError("Workbook header verification failed.")
    for row_number in range(2, sheet.max_row + 1):
        if sheet.cell(row_number, 4).value != "YES" or sheet.cell(row_number, 5).value != "YES":
            raise RuntimeError("Attendance values must be YES.")
        for column in range(1, 6):
            value = sheet.cell(row_number, column).value
            if isinstance(value, str) and value.startswith("="):
                raise RuntimeError("Formulas are not permitted in the upload file.")
    session.external_upload_generated_at = timezone.now()
    session.save(update_fields=["external_upload_generated_at", "updated_at"])
    ActivityLog.objects.create(
        actor=request.user,
        action=ActivityLog.Action.EXPORT,
        object_type="CourseSession",
        object_id=str(session.public_id),
        description="Exact external student upload workbook exported.",
        details={"student_rows": len(export_rows), "columns": HEADERS},
    )
    response = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        "attachment; filename*=UTF-8''"
        + quote(day_one_students_upload_filename(session))
    )
    return response
