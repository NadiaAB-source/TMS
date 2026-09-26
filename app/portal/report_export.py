"""Excel output for the very same filtered rows rendered by Reports."""
from datetime import date, datetime
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def workbook_bytes(title, headers, rows, filters=(), *, tz=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Report"
    sheet.append([title])
    sheet.append(list(headers))
    for values in rows:
        # Treat user-entered text as literal text, including strings starting =.
        # IDs remain strings; dates and counts remain real typed Excel values.
        sheet.append([None] * len(headers))
        for index, value in enumerate(values, 1):
            cell = sheet.cell(sheet.max_row, index)
            if isinstance(value, datetime) and value.tzinfo is not None:
                value = value.astimezone(tz).replace(tzinfo=None)
            if isinstance(value, str):
                value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                cell.value = value
                cell.data_type = "s"
            else:
                cell.value = value
            if isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(value, date):
                cell.number_format = "dd/mm/yyyy"
            if "Emirates ID" in str(headers[index - 1]):
                cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(sheet.max_row, 2)}"
    sheet.row_dimensions[1].height = 27
    sheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color="E9702B")
    for cell in sheet[2]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="383838")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 30
    for column in sheet.iter_cols(min_row=2):
        width = min(55, max(13, max(len(str(c.value or "")) for c in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column[1:]:
            cell.font = Font(name="Arial", size=11)
            if cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
    scope = book.create_sheet("Filters")
    scope.append(["Filter", "Value"])
    for label, value in filters:
        scope.append([str(label), str(value or "All")])
        for cell in scope[scope.max_row]:
            cell.data_type = "s"
    scope.column_dimensions["A"].width = 24
    scope.column_dimensions["B"].width = 65
    scope.freeze_panes = "A2"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def report_table(context):
    """Return headings and values without querying a second, unfiltered dataset."""
    view = context["report_view"]
    sessions = context["sessions"]
    if view == "progress":
        return ["Course", "Date from", "Date to", "Location", "Instructors", "Registration", "Status", "Enrolled", "Pass", "Fail", "Reject", "HP", "TTT"], [
            [s.display_course_name, s.start_date, s.end_date, str(s.camp or ""),
             "; ".join(str(i) for i in s.instructors.all()), "Open" if s.registration_published else "Closed", s.lifecycle_label,
             *[s.report_totals[k] for k in ("enrolled", "passed", "failed", "rejected", "hp", "ttt")]] for s in sessions]
    if view == "outcomes":
        return ["Course", "Date from", "Date to", "Location", "Enrolled", "Pass", "Fail", "Reject", "TBC", "HP", "TTT"], [
            [s.display_course_name, s.start_date, s.end_date, str(s.camp or ""),
             *[s.report_totals[k] for k in ("enrolled", "passed", "failed", "rejected", "pending", "hp", "ttt")]] for s in sessions]
    if view == "student_details":
        return ["Name in English", "Name in Arabic", "Emirates ID", "Email", "Course", "Date", "Location", "Status", "HP", "TTT"], [
            [r["name_english"], r["name_arabic"], r["formatted_eid"], r["email"], r["course_name"], r["date"],
             str(r["session"].camp or "") if r["session"] else "", r["outcome_label"], r["is_hp"], r["is_ttt"]] for r in context["student_rows"]]
    if view == "documents":
        selected = context["selected_report"]
        headers = ["Course", "Date", "Location"]
        if selected != "documents_frontend":
            headers += ["Day 1 list", "Stamped list", "Uploaded by", "Uploaded at"]
        if selected != "documents_stamped":
            headers += ["Front-End status", "Confirmed by", "Confirmed at"]
        rows = []
        for s in sessions:
            r = [s.display_course_name, s.start_date, str(s.camp or "")]
            stamped = s.current_stamped_list
            if selected != "documents_frontend":
                r += ["Saved" if s.latest_roster else "Awaiting generation", "Uploaded" if stamped else "Awaiting upload",
                      str(stamped.uploaded_by or "") if stamped else "", stamped.uploaded_at if stamped else None]
            if selected != "documents_stamped":
                r += [s.frontend_status, str(s.external_upload_confirmed_by or ""), s.external_upload_confirmed_at]
            rows.append(r)
        return headers, rows
    if view == "instructors":
        return ["Course", "Date", "Instructor", "Information sent", "Delivery"], [
            [a.session.display_course_name, a.session.start_date, a.instructor.name_english, a.notified_at,
             "Sent" if a.notified_at else "Not sent"] for a in context["instructor_rows"]]
    if view == "inventory":
        return ["Course", "Date", "Instructor", "Item", "Unit", "Used", "Notes", "Updated"], [
            [u.session.display_course_name, u.session.start_date, u.instructor.name_english,
             u.item.name, u.item.unit, u.quantity_used, u.notes, u.updated_at] for u in context["inventory_rows"]]
    if view == "duplicates":
        return ["Review status", "Course", "Date", "Name in English", "Name in Arabic", "Emirates ID", "Email", "Reason", "Reviewed by", "Reviewed at", "Review note"], [
            [r["review_label"], r["registration"].requested_session.course.title_english,
             r["registration"].requested_session.start_date, r["registration"].submitted_name_english,
             r["registration"].submitted_name_arabic, r["eid_display"], r["registration"].email_raw,
             "; ".join(f["title"] + ": " + f["detail"] for f in r["flags"]),
             str(r["registration"].duplicate_reviewed_by or ""), r["registration"].duplicate_reviewed_at,
             r["registration"].duplicate_review_notes] for r in context["duplicate_rows"]]
    return ["Date", "User", "Action", "Record type", "Record ID", "Description"], [
        [a.created_at, str(a.actor or "System"), a.get_action_display(), a.object_type,
         str(a.object_id), a.description] for a in context["recent_activity"]]
