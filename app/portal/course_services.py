from datetime import date
from pathlib import Path
import hashlib
import re

from django.conf import settings
from django.core.mail import EmailMessage
from django.db import transaction
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .identity import format_emirates_id
from .models import ActivityLog, CourseSequenceCounter
from .normalization import normalize_email
from .security import get_client_ip


STAMPED_TEMPLATE = (
    Path(__file__).resolve().parent
    / "resources"
    / "stamped_list_template.xlsx"
)
COURSE_SEQUENCE_COUNTER_KEY = "course_session"


def allocate_course_sequence_number():
    """Reserve the next permanent number when approved numbering is enabled.

    A locked singleton avoids duplicate numbers when two authorised users create
    courses at the same time.  The ``get_or_create`` fallback also keeps test
    databases and restored databases safe after a flush.
    """

    with transaction.atomic():
        CourseSequenceCounter.objects.get_or_create(
            key=COURSE_SEQUENCE_COUNTER_KEY,
            defaults={"next_number": 1},
        )
        counter = CourseSequenceCounter.objects.select_for_update().get(
            key=COURSE_SEQUENCE_COUNTER_KEY
        )
        if not counter.enabled:
            return None
        sequence_number = counter.next_number
        counter.next_number = sequence_number + 1
        counter.save(update_fields=["next_number"])
        return sequence_number


def enable_course_sequence_numbering(first_number=1):
    """Enable a freshly reset historical sequence at a deliberate number."""

    if first_number < 1:
        raise ValueError("The first course number must be at least 1.")
    with transaction.atomic():
        counter, _ = CourseSequenceCounter.objects.get_or_create(
            key=COURSE_SEQUENCE_COUNTER_KEY,
            defaults={"next_number": first_number},
        )
        counter = CourseSequenceCounter.objects.select_for_update().get(
            key=counter.key
        )
        counter.next_number = first_number
        counter.enabled = True
        counter.save(update_fields=["next_number", "enabled"])


def safe_reference_code(course, camp, start_date):
    base = course.code or course.title_english or "COURSE"
    base = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").upper()
    camp_part = re.sub(r"[^A-Za-z0-9]+", "-", camp.name).strip("-").upper()
    return f"{base}-{start_date:%Y%m%d}-{camp_part}"[:100]


def display_course_name(session):
    """Return the visible course name, including its permanent sequence."""
    course_name = (
        session.course.title_english
        or session.course.code
        or "Course"
    )
    course_name = re.sub(r"\s*[-–—]\s*", " · ", course_name).strip()
    if session.sequence_number:
        return f"{session.sequence_number:04d} · {course_name.upper()}"
    return course_name.upper()


def display_reference(session):
    """Return a short, human-readable reference without changing stored IDs."""
    course_name = display_course_name(session)
    date_text = (
        session.start_date.strftime("%d %b %Y").upper()
        if session.start_date
        else "DATE PENDING"
    )
    return f"{course_name.upper()} · {date_text}"


def course_date_text(session):
    if session.start_date and session.end_date:
        if session.start_date == session.end_date:
            return session.start_date.strftime("%d %B %Y")
        return f"{session.start_date:%d %B %Y} – {session.end_date:%d %B %Y}"
    return "Dates not set"


def _filename_component(value, fallback):
    """Return a readable filename part using the example's underscore style."""
    component = re.sub(r"[^\w]+", "_", str(value or "").strip(), flags=re.UNICODE)
    return component.strip("_") or fallback


def course_list_filename_base(session):
    """Build the shared Day 1 filename from dates, instructors and camp."""
    start = session.start_date
    end = session.end_date
    if start and end:
        if start == end:
            date_part = start.strftime("%d_%B")
        elif start.year == end.year and start.month == end.month:
            date_part = f"{start:%d}-{end:%d}_{start:%B}"
        else:
            date_part = f"{start:%d_%B}-{end:%d_%B}"
    elif start:
        date_part = start.strftime("%d_%B")
    else:
        date_part = "Dates"

    assignments = session.instructor_assignments.select_related(
        "instructor"
    ).order_by("id")
    first_names = []
    for assignment in assignments:
        full_name = (assignment.instructor.name_english or "").strip()
        first_name = full_name.split(maxsplit=1)[0] if full_name else "Instructor"
        first_names.append(_filename_component(first_name, "Instructor"))
    instructors_part = "-".join(first_names) or "Instructor"
    camp_part = _filename_component(
        session.camp.name if session.camp else "Camp",
        "Camp",
    )
    return f"{date_part}_{instructors_part}_{camp_part}"


def day_one_list_filename(session):
    return f"{course_list_filename_base(session)}.xlsx"


def day_one_students_upload_filename(session):
    return f"{course_list_filename_base(session)}_Students upload.xlsx"


def send_course_information(request, session, assignments):
    course_url = request.build_absolute_uri(
        reverse(
            "instructor_course_workspace",
            kwargs={"public_id": session.public_id},
        )
    )
    subject = f"Course assignment: {session.course.title_english}"
    body = "\n".join(
        [
            "You have been assigned to the following course:",
            "",
            f"Course: {session.course.title_english}",
            f"Reference: {session.reference_code}",
            f"Dates: {course_date_text(session)}",
            f"Camp: {session.camp}",
            f"Service Branch: {session.service_branch_name}",
            f"Capacity: {session.capacity}",
            f"Instructor-to-student ratio: 1:{session.instructor_student_ratio}",
            "",
            f"Point of Contact: {session.poc_name}",
            f"Contact number: {session.poc_contact_number}",
            f"Location: {session.poc_location_url}",
            "",
            "Log in and open the shared course workspace:",
            course_url,
        ]
    )
    sent_at = timezone.now()
    failures = []
    for assignment in assignments:
        instructor = assignment.instructor
        email = normalize_email(
            instructor.email
            or (instructor.user.email if instructor.user_id else "")
        )
        if not email:
            failures.append(f"{instructor.name_english}: delivery failed")
            continue
        try:
            delivered = EmailMessage(subject=subject, body=body, to=[email]).send(
                fail_silently=False
            )
        except Exception:
            failures.append(f"{instructor.name_english}: delivery failed")
            continue
        if delivered != 1:
            failures.append(f"{instructor.name_english}: delivery failed")
            continue
        assignment.notified_at = sent_at
        assignment.save(update_fields=["notified_at", "updated_at"])
    if not failures:
        session.course_information_sent_at = sent_at
        session.save(update_fields=["course_information_sent_at", "updated_at"])
    ActivityLog.objects.create(
        actor=request.user,
        action=ActivityLog.Action.UPDATE,
        object_type="CourseSession",
        object_id=str(session.public_id),
        description=(
            "Course information delivered to all assigned instructors."
            if not failures
            else "Course information delivery incomplete."
        ),
        details={
            "assigned_count": len(assignments),
            "delivered_count": len(assignments) - len(failures),
            "failed_instructors": [failure.split(":", 1)[0] for failure in failures],
        },
        ip_address=get_client_ip(request),
    )
    return failures


def _sheet_header(session):
    if session.start_date and session.end_date:
        if session.start_date == session.end_date:
            dates = session.start_date.strftime("%d %B")
        else:
            dates = f"{session.start_date:%d}-{session.end_date:%d %B}"
    else:
        dates = "Dates"
    assignments = session.instructor_assignments.select_related(
        "instructor"
    ).order_by("instructor__name_english")
    instructor_names = ", ".join(
        assignment.instructor.name_english for assignment in assignments
    ) or "Instructor"
    camp = session.camp.name if session.camp else "Camp"
    return f"{dates}_{instructor_names}_{camp}_Stamped_List"


def _remove_unused_student_rows(sheet, student_count):
    """Keep the template footer while removing unneeded numbered rows."""
    unused_count = 20 - student_count
    if unused_count <= 0:
        return 26

    first_unused_row = 4 + student_count
    last_unused_row = 23
    merged_ranges = [str(item) for item in sheet.merged_cells.ranges]
    for merged_range in merged_ranges:
        sheet.unmerge_cells(merged_range)

    sheet.delete_rows(first_unused_row, unused_count)

    for merged_range in merged_ranges:
        min_column, min_row, max_column, max_row = range_boundaries(
            merged_range
        )
        if min_row >= 24:
            min_row -= unused_count
            max_row -= unused_count
        elif max_row >= first_unused_row:
            raise RuntimeError(
                "The stamped-list template has a merge inside student rows."
            )
        sheet.merge_cells(
            start_row=min_row,
            start_column=min_column,
            end_row=max_row,
            end_column=max_column,
        )

    return 26 - unused_count


def build_stamped_roster_workbook(session, registrations):
    if not STAMPED_TEMPLATE.is_file():
        raise RuntimeError("The stamped-list template is missing.")
    registrations = list(registrations)
    page_count = max(1, (len(registrations) + 19) // 20)
    workbook = load_workbook(STAMPED_TEMPLATE)
    source = workbook.worksheets[0]
    sheets = [source]
    for _ in range(1, page_count):
        sheets.append(workbook.copy_worksheet(source))

    for page_index, sheet in enumerate(sheets):
        start = page_index * 20
        page_rows = registrations[start:start + 20]
        last_serial = start + len(page_rows)
        sheet.title = f"List {start + 1}-{last_serial}"
        sheet["A1"] = f"Course: {session.course.title_english}"
        sheet["C1"] = _sheet_header(session)
        print_end_row = _remove_unused_student_rows(sheet, len(page_rows))
        for row_number in range(4, 4 + len(page_rows)):
            for column_number in range(2, 9):
                sheet.cell(row=row_number, column=column_number).value = None
        for offset, registration in enumerate(page_rows):
            row_number = 4 + offset
            student = registration.student
            eid = (
                (student.eid if student else "")
                or registration.eid_normalized
                or registration.eid_raw
            )
            eid = format_emirates_id(eid, fallback="")
            sheet.cell(row=row_number, column=1).value = start + offset + 1
            sheet.cell(row=row_number, column=2).value = (
                registration.email_raw or (student.email if student else "")
            )
            sheet.cell(row=row_number, column=3).value = eid
            sheet.cell(row=row_number, column=4).value = (
                registration.submitted_name_arabic
                or (student.name_arabic if student else "")
            )
            sheet.cell(row=row_number, column=5).value = (
                registration.submitted_name_english
                or (student.name_english if student else "")
            ).upper()
        sheet.print_area = f"A1:H{print_end_row}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1

    from io import BytesIO

    output = BytesIO()
    workbook.save(output)
    data = output.getvalue()
    verification = load_workbook(BytesIO(data), data_only=False)
    if len(verification.worksheets) != page_count:
        raise RuntimeError("Saved roster page verification failed.")
    for sheet in verification.worksheets:
        printed_end_row = int(str(sheet.print_area).rsplit("$", 1)[-1])
        student_rows = printed_end_row - 6
        for row_number in range(4, 4 + student_rows):
            if sheet.cell(row=row_number, column=1).value in {None, 0, "0"}:
                raise RuntimeError(
                    "Saved roster contains an empty or zero student row."
                )
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise RuntimeError("Formulas are not permitted in saved rosters.")
    return data, hashlib.sha256(data).hexdigest()
